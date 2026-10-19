"""Traceable local readers. Raw inputs are never overwritten."""
from pathlib import Path
import hashlib
import numpy as np
import pandas as pd
from .schema import INPUT, ExperimentConfig, supsi_spec
from .physics import thermal_conditions


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def laboratory_from_xlsx(path=None):
    from openpyxl import load_workbook
    path = path or INPUT / "SUPSI/raw/laboratory/IEC61853_1-2.xlsx"
    sheet = load_workbook(path, read_only=True, data_only=True)["61853-1 & 2"]
    rows = []
    # Only these positions are absent in this IEC matrix; do NOT discard all 32s.
    missing = {(1100,15),(400,75),(200,50),(200,75),(100,50),(100,75)}
    for variable, first, unit in [("Pmp",9,"W"),("Isc",19,"A"),("Voc",29,"V")]:
        for row in range(first,first+7):
            g = float(sheet.cell(row,1).value)
            for col,t in enumerate([15,25,50,75],2):
                value = sheet.cell(row,col).value
                empty = value is None or isinstance(value,str) and not value.strip()
                if (g,t) in missing:
                    if not empty and str(value).strip() != "32":
                        raise ValueError(f"Matriz alterada: inspecione {row},{col}.")
                    continue
                if empty:
                    raise ValueError(f"Medição esperada ausente: {row},{col}.")
                rows.append(dict(variable=variable, irradiance_poa_Wm2=g,
                    temperature_module_C=float(t), measured_value=float(value), unit=unit,
                    xlsx_row=row, xlsx_col=sheet.cell(row,col).column_letter))
    out = pd.DataFrame(rows)
    if out.groupby("variable").size().to_dict() != {"Isc":22,"Pmp":22,"Voc":22}:
        raise ValueError("Matriz deve conter 22 × 3 observações.")
    return out


def laboratory():
    return pd.read_csv(INPUT / "SUPSI/derived/lab_IEC61853_matrix_long.csv")


def read_field(paths):
    frames = []
    for path in paths:
        path = Path(path)
        d = pd.read_csv(path,sep=";",encoding="utf-8-sig")
        d.columns = d.columns.str.strip()
        required = {"Date","Time","Pm","Tbom","Gpoa","AIR_TEMP","WIND_SPEED"}
        if not required.issubset(d.columns):
            raise ValueError(f"{path.name}: colunas ausentes: {required-set(d.columns)}")
        d.insert(0,"source_row",np.arange(2,len(d)+2))
        d.insert(0,"source_file",path.name)
        d["timestamp"] = pd.to_datetime(d.Date.str.strip()+" "+d.Time.str.strip(),
            format="%d.%m.%Y %H:%M:%S",errors="coerce")
        for col in required-{"Date","Time"}:
            d[col] = pd.to_numeric(d[col],errors="coerce")
        frames.append(d)
    if not frames:
        raise ValueError("Nenhum CSV fornecido.")
    d = pd.concat(frames,ignore_index=True)
    d.insert(0,"source_order",np.arange(len(d)))
    return d.sort_values(["timestamp","source_order"],kind="stable").reset_index(drop=True)


def field_paths(months):
    return [INPUT/"SUPSI/raw/field_2019"/(month+".csv") for month in months]


def quality_flags(raw, config: ExperimentConfig):
    d = raw.copy()
    reasons = [[] for _ in range(len(d))]
    def add(mask, reason):
        for i in np.flatnonzero(np.asarray(mask)):
            reasons[i].append(reason)
    add(d.timestamp.isna(),"invalid_timestamp")
    add(d.timestamp.duplicated(keep=False),"duplicate_timestamp")
    for col in ["Pm","Gpoa","Tbom"] + (["AIR_TEMP","WIND_SPEED"] if config.thermal=="faiman" else []):
        add(~np.isfinite(d[col]),"nonfinite_"+col)
    add(d.Pm < 0,"negative_Pm")
    add(d.Gpoa < 0,"negative_Gpoa")
    add((d.Gpoa>=0) & (d.Gpoa<config.g_min_Wm2),"below_G_threshold")
    add(d.Gpoa > config.g_max_Wm2,"above_G_limit")
    add((d.Tbom<config.temperature_min_C)|(d.Tbom>config.temperature_max_C),"implausible_Tbom")
    if config.thermal=="faiman":
        add(d.WIND_SPEED<0,"negative_wind")
        add((d.AIR_TEMP < -60)|(d.AIR_TEMP > 60),"implausible_Tamb")
    d["eligible"] = [not r for r in reasons]
    d["exclusion_reasons"] = [";".join(r) for r in reasons]
    d["lag_minutes"] = d.timestamp.diff().dt.total_seconds()/60
    d["gap_before"] = d.lag_minutes > config.max_gap_minutes
    # Context only; outside calibration domain is reported, NOT silently excluded.
    d["outside_lab_rectangle"] = ~d.Gpoa.between(100,1100) | ~d.Tbom.between(15,75)
    return d


def prepared_conditions(quality, config, spec=None):
    d = quality.loc[quality.eligible].copy()
    if d.empty:
        raise ValueError("Nenhuma amostra elegível. Verifique os filtros e as fontes.")
    g = d.Gpoa.to_numpy(float)
    tm,tc = thermal_conditions(g, d.Tbom.to_numpy(float), d.AIR_TEMP.to_numpy(float),
        d.WIND_SPEED.to_numpy(float), method=config.thermal,delta_T_C=config.delta_T_C,
        spec=spec or supsi_spec())
    return d, {"g":g,"tm":tm,"tc":tc}


def custom_profile(raw, timestamp_col, irradiance_col, temperature_col, *, irradiance_kind,
                   timestamp_format, temperature_kind, spec, delta_T_C=3.0):
    if irradiance_kind != "POA":
        raise ValueError("GHI não é POA. Forneça irradiância no plano; transposição não é automática.")
    if not timestamp_format.strip():
        raise ValueError("Declare o formato de data/hora, por exemplo %Y-%m-%d %H:%M:%S.")
    times = pd.to_datetime(raw[timestamp_col],format=timestamp_format,errors="coerce")
    g = pd.to_numeric(raw[irradiance_col],errors="coerce").to_numpy(float)
    t = pd.to_numeric(raw[temperature_col],errors="coerce").to_numpy(float)
    if times.isna().any() or times.duplicated().any() or not times.is_monotonic_increasing:
        raise ValueError("Datas inválidas, duplicadas ou fora de ordem; corrija explicitamente o CSV.")
    if not np.all(np.isfinite(g)) or not np.all(np.isfinite(t)) or np.any(g<0):
        raise ValueError("CSV contém irradiância/temperatura inválida; nenhuma linha foi removida.")
    method = {"Tmod":"measured_back","Tcell":"measured_cell","Tamb":"noct"}[temperature_kind]
    tm,tc = thermal_conditions(g,t_back=t,t_cell=t,t_ambient=t,method=method,spec=spec,delta_T_C=delta_T_C)
    return pd.DataFrame({"timestamp":times,"Gpoa":g,"Tmod_C":tm,"Tcell_C":tc})
